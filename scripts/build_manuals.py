# -*- coding: utf-8 -*-
"""操作説明書 (お客様向け / 社内営業向け) を python-docx で生成する
    python build_manuals.py <出力フォルダ> <スクリーンショットフォルダ> [GitHubユーザー名]
"""
import os, sys, datetime
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = sys.argv[1] if len(sys.argv) > 1 else "."
SHOTS = sys.argv[2] if len(sys.argv) > 2 else "."
GH = sys.argv[3] if len(sys.argv) > 3 else "＜GitHubユーザー名または組織名＞"
SITE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "site")
TODAY = datetime.date.today().strftime("%Y年%m月%d日")
REPO = f"https://github.com/{GH}/bimcim-lod-viewer"
PAGES = f"https://{GH}.github.io/bimcim-lod-viewer/"
RELEASES = f"{REPO}/releases"
ARTIFACT = "https://claude.ai/artifact/QZ7fGMTNZt6QEi8HKHEZGS"

def set_font(run, size=None, bold=None, color=None):
    run.font.name = "Yu Gothic"; run._element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic")
    if size: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if color: run.font.color.rgb = RGBColor(*color)

def new_doc(title, subtitle):
    d = Document()
    for s in d.sections:
        s.page_width, s.page_height = Mm(210), Mm(297); s.left_margin = s.right_margin = Mm(20); s.top_margin = s.bottom_margin = Mm(20)
    st = d.styles["Normal"]; st.font.name = "Yu Gothic"; st.element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic"); st.font.size = Pt(10.5)
    for lv, sz in ((1, 16), (2, 13), (3, 11.5)):
        hs = d.styles[f"Heading {lv}"]; hs.font.name = "Yu Gothic"; hs.element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic"); hs.font.size = Pt(sz); hs.font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F)
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(120)
    r = p.add_run(title); set_font(r, 24, True, (0x1E, 0x3A, 0x5F))
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = p.add_run(subtitle); set_font(r, 13, False, (0x50, 0x5A, 0x64))
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(200)
    r = p.add_run(f"Malme株式会社\n{TODAY}  第1.0版"); set_font(r, 11)
    d.add_page_break()
    return d

def H(d, text, lv=1): d.add_heading(text, lv)
def P(d, text, bold=False, size=None, color=None, italic=False):
    p = d.add_paragraph(); r = p.add_run(text); set_font(r, size, bold, color); r.italic = italic; return p
def B(d, items, style="List Bullet"):
    for it in items:
        p = d.add_paragraph(style=style); r = p.add_run(it); set_font(r)
def N(d, items):
    """番号付きリスト (Word の自動番号は前のリストから連番になるため手動で番号を振る)"""
    for i, it in enumerate(items, 1):
        p = d.add_paragraph(); p.paragraph_format.left_indent = Mm(8); p.paragraph_format.first_line_indent = Mm(-6)
        r = p.add_run(f"{i}. {it}"); set_font(r)
def NOTE(d, text):
    t = d.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER; c = t.rows[0].cells[0]
    shade(c, "FFF4E5"); p = c.paragraphs[0]; r = p.add_run(text); set_font(r, 10)
    d.add_paragraph()
def shade(cell, hexcolor):
    tcPr = cell._element.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexcolor); tcPr.append(shd)
def TABLE(d, header, rows, widths=None):
    t = d.add_table(rows=1, cols=len(header)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; shade(c, "DCE6F1"); r = c.paragraphs[0].add_run(h); set_font(r, 9.5, True)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            r = cells[i].paragraphs[0].add_run(str(v)); set_font(r, 9.5)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Mm(w)
    d.add_paragraph()
def IMG(d, path, width_mm=160, caption=None):
    if not os.path.exists(path): P(d, f"（画像: {os.path.basename(path)} — 準備中）", italic=True, color=(0x99, 0x99, 0x99)); return
    d.add_picture(path, width=Mm(width_mm)); d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if caption:
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = p.add_run(caption); set_font(r, 9, False, (0x60, 0x60, 0x60))

MODELS = [
    ("山岳トンネル（NATM 2車線）", "トンネル", "延長240m。支保パターン DⅢa/DⅡ/DⅠ、坑門工、非常駐車帯、ロックボルト・鋼製支保工・覆工鉄筋", ""),
    ("道路土工（切土・盛土）", "土工", "2車線道路240m。切土1:1.0/盛土1:1.5、小段、法尻側溝、舗装4層、法面保護工、縦排水、ガードレール", ""),
    ("PC橋（ポステンT桁 2径間）", "橋梁", "橋長61m。T桁5本、逆T式橋台・壁式橋脚・杭、支承・伸縮装置・高欄、PC鋼材・配筋", ""),
    ("BOXカルバート（1連 3.0×2.5m）", "カルバート", "延長40m。頂版・側壁・底版・ハンチ、ウイング、取付水路、目地・止水板、配筋", "試作中"),
    ("擁壁（逆T式 H=5m＋重力式 H=2m）", "擁壁", "延長40m。基礎、目地・止水板、水抜き孔、裏込め・透水マット、配筋、防護柵", ""),
    ("電線共同溝（管路部・特殊部）", "電線共同溝", "管路部240m。電力6条・通信 共用FA方式、電力/通信ハンドホール、引込管、横断管、地上機器", ""),
    ("法枠工・鉄筋挿入工・かご工", "法面工・護岸", "切土法面60m。吹付法枠300×300@2.0m、主アンカー、鉄筋挿入工、排水孔、ふとんかご3段", ""),
]
LOD_ROWS = [
    ("100", "位置", "対象を記号や線、単純な形状でその位置を示したモデル。"),
    ("200", "構造形式", "対象の構造形式が分かる程度のモデル。標準横断面を対象範囲でスイープさせて作成する程度の表現。"),
    ("300", "外形形状", "附帯工等の細部構造、接続部構造を除き、対象の外形形状を正確に表現したモデル。"),
    ("400", "細部・配筋", "詳細度300に加えて、附帯工、接続構造などの細部構造及び配筋も含めて、正確にモデル化する。"),
    ("500", "完成形状", "対象の現実の形状を表現したモデル。"),
]
DOCK_ROWS = [
    ("≡", "折りたたみ", "アイコン列を畳む／開く"),
    ("カメラ", "視点", "全景・坑口・内部・断面など、モデルごとに用意した視点へ切替"),
    ("目", "表示", "ワイヤーフレーム、覆工（主構造）の透過、回転中心の目印、回転方向（AutoCAD風/反転）"),
    ("立方体に切断線", "断面カット", "任意の測点で切断表示。反転で反対側を表示"),
    ("メーター", "感度", "回転感度・ズーム感度を各6段階で調整（右下の数字が現在の回転感度）"),
    ("中心点つき円弧矢印", "ピボット", "回転半径の決め方：固定／自動／1〜150m（右下に現在値）"),
    ("十字", "回転中心を置く", "クリックで指定、測点に置く、モデル中央、中心へ寄る"),
    ("虫めがね", "クリックした点へ寄る", "押してからモデル上の点をクリックすると、その点を中心に視点が寄る（繰り返すと段階的に接近）"),
]

def common_operation_sections(d, brief=False):
    H(d, "画面の見方")
    IMG(d, os.path.join(SHOTS, "index.png"), 160, "図：入口画面（工種を選ぶ）")
    IMG(d, os.path.join(SHOTS, "viewer.png"), 160, "図：ビューア画面（PC橋 詳細度300）")
    TABLE(d, ["位置", "名称", "役割"], [
        ("上部", "タイトル・詳細度ボタン", "選択中の工種名と概要。詳細度100〜500のボタンでモデルを切替。「工種を選ぶ」で入口画面に戻る"),
        ("左端", "ツールドック", "視点・表示・断面カット・感度・ピボット・回転中心・寄る の各機能（下表）"),
        ("右上", "ViewCube", "キューブをドラッグで回転、面・辺・角をクリックでその方向からの視点。下の家アイコンでホーム視点"),
        ("右", "情報パネル", "選択中の詳細度の定義、部材数・三角形数、部材一覧（クリックで表示/非表示）、モデルの前提・注記"),
        ("左下", "回転中心の表示", "回転中心の測点・高さ・視点までの距離"),
    ], [18, 32, 120])
    H(d, "基本操作（マウス・タッチ）")
    TABLE(d, ["操作", "マウス（PC）", "タッチ（スマホ・タブレット）"], [
        ("回転", "左ドラッグ", "1本指でドラッグ"),
        ("移動（パン）", "右ドラッグ", "2本指でドラッグ"),
        ("拡大縮小", "ホイール（カーソル位置に向かって寄る）", "2本指でピンチ"),
        ("回転中心を置く", "ダブルクリック", "ダブルタップ"),
        ("指定した点へ寄る", "虫めがねアイコン→クリック（Shift＋ダブルクリックでも可）", "虫めがねアイコン→タップ"),
    ], [30, 65, 75])
    P(d, "回転方向は AutoCAD と同じ「ドラッグした方向にモデルが回る」向きです。表示パネルの「回転: AutoCAD風」を OFF にすると逆向きになります。")
    H(d, "左端ツールドックの機能")
    TABLE(d, ["アイコン", "機能", "内容"], DOCK_ROWS, [30, 35, 105])
    H(d, "詳細度の切替と情報パネル")
    B(d, ["上部の 100 / 200 / 300 / 400 / 500 ボタンで切り替えます。詳細度ごとに別のモデルを読み込むため、初回は数秒かかります（詳細度400・500は部材が多く、最大10MB程度）。",
          "右パネル上部にその詳細度の「共通定義」と「この工種でのモデル化」の説明が表示されます。",
          "「モデルに含まれる要素」は分類ごとの部材一覧です。チェックを外すと非表示にできます。「詳細度○○で追加」のタグは、その詳細度で初めて追加された部材を示します。",
          "詳細度300では支保パターンの範囲や箱抜き位置などを半透明の「記号」で示し、詳細度400で実形状（ロックボルト・配筋など）に置き換わります。この違いを比べると詳細度の差が分かりやすくなります。"])
    if not brief:
        H(d, "断面カット")
        B(d, ["ツールドックの断面カットで「カットする」にチェックし、測点スライダーで切断位置を選びます。「反転」で反対側を表示します。",
              "視点の「断面」ボタンを押すと、あらかじめ設定した測点で切断した状態の視点になります。"])
        IMG(d, os.path.join(SITE, "models", "tunnel", "view_LOD400.png"), 140, "図：山岳トンネル 詳細度400（ロックボルト・鋼製支保工・防水シートが実形状）")
    H(d, "視点が分からなくなったとき")
    N(d, ["右上の家アイコン（ホーム）を押すと全景に戻り、回転中心がモデル中央に置き直されます。",
          "回転中心アイコン（十字）→「モデル中央」→「中心へ寄る」で、モデル中央の30m手前に視点が移動します。",
          "ViewCube の「上」をクリックすると平面図のように真上から（北が画面上）見えます。",
          "ホイールで寄りすぎたときは、ズーム感度を下げる（感度アイコン→ズーム 1〜2）と細かく合わせられます。"])

def customer_manual():
    d = new_doc("BIM/CIM 詳細度サンプルビューア\n操作説明書", "工種別サンプル3Dモデルで詳細度100〜500を確認するための手引き")
    H(d, "1. はじめに")
    P(d, "本ビューアは、国土交通省「BIM/CIM活用ガイドライン（案）第1編 共通編」で定義されたBIM/CIMモデルの詳細度（100・200・300・400・500）を、工種ごとの汎用サンプル3Dモデルで実際に切り替えて比較できるWebアプリケーションです。事前協議で「どの詳細度でどこまで作り込むか」を具体的なモデルを見ながら確認する用途を想定しています。")
    NOTE(d, "収録モデルはすべて架空の寸法・線形・地形で作成した汎用サンプルです。実在する案件の設計データは含まれていません。また各モデルの寸法・数量は一般的な値であり、設計の根拠として使用することはできません。")
    H(d, "2. 利用方法の選択")
    TABLE(d, ["利用方法", "こんなときに", "必要なもの"], [
        ("A. Web版（推奨）", "PC・スマホ・タブレットからすぐ見たい。社内で共有したい", "インターネット接続、最新のブラウザ（Chrome / Edge / Safari）"),
        ("B. オフライン版（1ファイル）", "現場や会議室などネット接続のない環境で使う。PCに保存して使う", "PC（Windows / Mac）、HTMLファイル約50MB"),
        ("C. モデルデータ", "自社のCIMソフト（Navisworks・AutoCAD 等）で開きたい", "IFC / DXF / OBJ を扱えるソフト"),
    ], [42, 68, 60])
    H(d, "2.1 Web版で開く", 2)
    B(d, [f"ブラウザで次のURLを開きます：{PAGES}",
          f"特定の工種を直接開くには URL の末尾に #ID を付けます（例：{PAGES}#tunnel）。ID は「6. 収録モデル一覧」を参照してください。",
          "スマホ・タブレットでも同じURLで動作します。画面が縦長のときは、3D表示の下に情報パネルが並びます。"])
    H(d, "2.2 GitHub からダウンロードする（オフライン版・モデルデータ）", 2)
    N(d, [f"ブラウザで公開ページを開きます：{RELEASES}",
          "最新のリリース（一番上）の「Assets」を展開します。",
          "必要なファイルをクリックしてダウンロードします。",
          "ブラウザによっては「このファイルは一般的にダウンロードされていません」等の警告が出ることがあります。発信元が本リポジトリであることを確認のうえ「保存」を選んでください。",
          "ダウンロードしたファイルは任意のフォルダ（例：デスクトップ）に保存します。zip の場合は右クリック→「すべて展開」で解凍してください。"])
    TABLE(d, ["ファイル名", "内容", "使い方"], [
        ("bimcim-lod-viewer_offline.html", "オフライン版ビューア（全7工種・全詳細度を1ファイルに同梱、約50MB）", "ダブルクリックでブラウザが開き、そのまま使えます。インターネット接続は不要です"),
        ("＜工種ID＞_ifc_dxf_obj.zip", "各工種の IFC4 / DXF（3Dメッシュ）/ OBJ ファイル（詳細度100〜500）", "解凍して、お使いのCIMソフトで開きます。IFCの各部材には詳細度などの属性が付いています"),
        ("Source code (zip)", "ビューアとモデル生成スクリプトのソース一式", "開発者向け。ローカルサーバで動かす場合や、モデルを再生成する場合に使用"),
    ], [55, 60, 55])
    NOTE(d, "オフライン版HTMLを開いてもモデルが表示されない場合は、ブラウザが古いか WebGL が無効になっている可能性があります。Chrome または Edge の最新版でお試しください。ファイルサイズが大きいため、開くまでに数秒〜十数秒かかります。")
    H(d, "2.3 ソース一式をダウンロードしてローカルで動かす（開発者向け）", 2)
    N(d, [f"リポジトリのトップページ（{REPO}）で緑色の「Code」ボタン→「Download ZIP」を選び、解凍します。",
          "index.html はブラウザの制約により直接ダブルクリックでは動きません（オフライン版HTMLは動きます）。解凍したフォルダでコマンドプロンプトを開き、簡易サーバを起動します（Python がある場合）： python -m http.server 8000",
          "ブラウザで http://localhost:8000/ を開きます。"])
    H(d, "3. 動作環境")
    TABLE(d, ["項目", "推奨"], [
        ("ブラウザ", "Google Chrome / Microsoft Edge / Safari の最新版（WebGL 対応）"),
        ("PC", "Windows 10/11、macOS。メモリ 8GB 以上。詳細度400・500は三角形数が多いため、内蔵GPUのPCでは動作が重くなることがあります"),
        ("スマホ・タブレット", "iOS 16 以降の Safari、Android の Chrome。Web版の利用を推奨（オフライン版は約50MBのため PC 推奨）"),
        ("通信", "Web版はモデル読込時に通信します（詳細度400・500で最大10MB程度）。オフライン版は不要"),
    ], [40, 135])
    common_operation_sections(d)
    H(d, "6. 収録モデル一覧")
    TABLE(d, ["ID", "工種", "内容", "備考"], [(("tunnel", "earthwork", "pcbridge", "boxculvert", "retwall", "conduit", "slope")[i], m[0], m[2], m[3]) for i, m in enumerate(MODELS)], [22, 40, 95, 18])
    NOTE(d, "「試作中」と表示されるモデルは形状が暫定で、実際の構造物とは細部が異なります。詳細度の説明用としてご覧ください。")
    H(d, "7. 詳細度の定義（共通定義）")
    TABLE(d, ["詳細度", "略称", "共通定義（BIM/CIM活用ガイドライン（案）共通編 表1-2）"], LOD_ROWS, [18, 22, 135])
    P(d, "山岳トンネルの工種別の定義文は同表の「山岳トンネルの例」の記載を用いています。その他の工種の定義文は共通定義に基づく本サンプルでの適用例であり、協議にあたっては『3次元モデル成果物作成要領（案）』および各分野編の記載を正としてください。")
    H(d, "8. よくある質問")
    TABLE(d, ["質問", "回答"], [
        ("モデルの寸法は実際の設計に使えますか", "使えません。一般的な値で作成した架空のサンプルです。"),
        ("自社の案件のモデルで同じように詳細度を確認できますか", "生成スクリプト（Python）のパラメータを変えることで別条件のサンプルを作成できます。個別案件への適用はご相談ください。"),
        ("スマホで詳細度400を開くと重い", "詳細度400・500は部材が多いため、スマホでは300までの確認を推奨します。"),
        ("IFC を Navisworks で開くと色が付かない", "IFC4 の表面スタイルを付与しています。表示されない場合は DXF をお試しください。"),
        ("印刷したい", "ブラウザの印刷機能ではなく、画面のスクリーンショットをご利用ください。"),
    ], [70, 105])
    H(d, "9. ライセンス・出典")
    B(d, ["ビューアおよびモデルは MIT ライセンスで公開しています。社内利用・改変・再配布は自由ですが、無保証です。",
          "詳細度の定義文の出典：国土交通省「BIM/CIM活用ガイドライン（案）第1編 共通編」（令和4年3月）表1-2、社会基盤情報標準化委員会「土木分野におけるモデル詳細度標準（案）【改訂版】」（平成30年3月）",
          "3D描画ライブラリ：three.js（MIT ライセンス）"])
    H(d, "10. お問い合わせ")
    P(d, "Malme株式会社　担当：＿＿＿＿＿＿＿　E-mail：＿＿＿＿＿＿＿＿＿＿")
    return d

def sales_manual():
    d = new_doc("BIM/CIM 詳細度サンプルビューア\n社内向け操作説明書", "営業・CS向け：アーティファクト共有版の使い方と説明の流れ")
    H(d, "1. このビューアの使いどころ")
    B(d, ["発注者・受注者の事前協議で「詳細度○○とはどこまで作り込むか」を、言葉や静止画ではなく3Dモデルの切替で見せるためのツールです。",
          "提案時に「当社はこのレベルまで作れます」「この詳細度なら成果物はこうなります」を示す営業資料としても使えます。",
          "すべて架空寸法の汎用モデルなので、契約外のお客様に見せても問題ありません（実案件データは一切含みません）。"])
    H(d, "2. 開き方と共有方法")
    H(d, "2.1 社内（アーティファクト版）", 2)
    B(d, [f"URL：{ARTIFACT}",
          "claude.ai にログインしたブラウザで開きます。スマホ・タブレットでも開けます。",
          "アーティファクトの共有設定は、ページ右上の「Share（共有）」メニューから所有者が変更します。社内メンバーに共有する場合は所有者に依頼してください。",
          "画面上部の「工種を選ぶ」で入口画面に戻り、別の工種を選べます。URL 末尾に #tunnel のように付けると直接その工種を開けます。"])
    H(d, "2.2 社外（お客様）に見せる・渡す場合", 2)
    TABLE(d, ["方法", "URL / ファイル", "備考"], [
        ("Web版（GitHub Pages）", PAGES, "お客様のPC・スマホからそのまま閲覧可。案内は「お客様向け操作説明書」を添付"),
        ("オフライン版HTML", f"{RELEASES} の bimcim-lod-viewer_offline.html", "1ファイル約50MB。ネット接続不要。会議室・現場での説明向き"),
        ("モデルデータ（IFC/DXF/OBJ）", f"{RELEASES} の ＜工種ID＞_ifc_dxf_obj.zip", "お客様のCIMソフトで開いてもらう場合"),
    ], [45, 75, 55])
    NOTE(d, "アーティファクト版の URL は社外には共有しないでください（claude.ai のアカウントが必要で、社内共有向けです）。社外向けには上表の Web版 / オフライン版を案内します。")
    common_operation_sections(d, brief=True)
    H(d, "説明の流れ（デモの進め方・約10分）")
    N(d, ["入口画面でお客様の工種に近いモデルを選ぶ（迷ったら山岳トンネル：定義文がガイドラインの記載そのままで説明しやすい）。",
          "詳細度100 → 200 の順に切り替え、「位置だけ」「標準断面をスイープしただけ」の違いを見せる。",
          "詳細度300で、支保パターン範囲や箱抜き位置が半透明の「記号」で示されていることを指摘する（この段階では細部は作らない）。",
          "詳細度400に切り替え、記号が消えてロックボルト・配筋・支保工の実形状になることを見せる。右パネルの「部材数・三角形数」の増え方で作業量の差を説明する。",
          "断面カット（視点の「断面」ボタン）で内部構造を見せる。ホイールで寄り、ダブルクリックで回転中心を置いて細部を回して見せる。",
          "詳細度500は「完成形状（出来形・完成設備）」の位置づけであることを説明する。",
          "最後に「詳細度の定義は共通編の表1-2、工種別は成果物作成要領・各分野編が正」であることを補足する。"])
    H(d, "各工種の見せどころ")
    TABLE(d, ["工種", "見せどころ"], [
        ("山岳トンネル", "300の支保パターン範囲記号 → 400のロックボルト全数・鋼製支保工。坑門工と非常駐車帯の拡幅"),
        ("道路土工", "200の固定断面と300の地形擦り付け（法尻の位置が変わる）。400の吹付法枠・植生工・排水"),
        ("PC橋", "200の概形（桁・欄干・橋台）→ 300の桁形状・杭 → 400のPC鋼材（放物線）・床版配筋・支承"),
        ("擁壁", "300の水抜き孔・目地「記号」→ 400の配筋・裏込め・止水板"),
        ("電線共同溝", "「表示」で路面が半透明。300の管路一体形状 → 400の管条ごとの管路・ハンドホール配筋・引込管"),
        ("法枠工・かご工", "300の法枠範囲パターン → 400の格子・鉄筋挿入工・中詰石。かご工の段ごとの表現"),
        ("BOXカルバート", "【試作中】形状は暫定。詳細度の考え方（目地記号→止水板、配筋）の説明のみに使う"),
    ], [40, 135])
    H(d, "よく聞かれる質問と答え方")
    TABLE(d, ["質問", "答え方"], [
        ("これは実案件のモデル？", "いいえ。すべて架空寸法の汎用サンプルで、実案件データは含まれていません（README と画面上に明記）。"),
        ("詳細度の定義はどこから？", "国交省 BIM/CIM活用ガイドライン（案）共通編 表1-2。山岳トンネルは同表の記載そのまま、他工種は共通定義に基づく当社の適用例。"),
        ("うちの案件でも作れる？", "生成スクリプト（Python）のパラメータで断面・延長を変えて再生成できる。個別案件は別途相談。"),
        ("詳細度400と300で作業量はどれくらい違う？", "画面の部材数・三角形数を見せながら「400は配筋・ボルト等を全数作るため部材数が数倍になる」と説明。具体的な工数は案件ごとに見積り。"),
        ("スマホで重い", "詳細度400・500は重いので、スマホでは300までを推奨。PC版またはオフライン版を案内。"),
    ], [60, 115])
    H(d, "困ったとき・フィードバック")
    B(d, ["表示が崩れる・動かない：ブラウザを最新の Chrome / Edge にして再読込。それでも直らない場合はスクリーンショットと工種・詳細度を添えて開発担当へ。",
          "モデルの形状や定義文の修正要望：工種名・詳細度・修正内容を開発担当へ。生成スクリプトで再出力します。",
          "URL・ファイルの最新版：GitHub の Releases ページが最新です。"])
    return d

os.makedirs(OUT, exist_ok=True)
customer_manual().save(os.path.join(OUT, "BIMCIM詳細度サンプルビューア_操作説明書（お客様向け）.docx"))
sales_manual().save(os.path.join(OUT, "BIMCIM詳細度サンプルビューア_操作説明書（社内営業向け）.docx"))
print("written to", os.path.abspath(OUT))
